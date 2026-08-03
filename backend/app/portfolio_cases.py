from __future__ import annotations

import secrets
from pathlib import Path

import fitz
from docx import Document as WordDocument
from openpyxl import Workbook

from .config import settings
from .database import connect, json_dump, now_iso
from .facts import FACT_REGISTRY
from .parsers import sha256


def is_demo_project(project_id: str) -> bool:
    return project_id.startswith("demo-")


def _docx(path: Path, title: str, paragraphs: list[str]) -> None:
    document = WordDocument()
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def _xlsx(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, data in sheets.items():
        sheet = workbook.create_sheet(name)
        for row_values in data:
            sheet.append(row_values)
    workbook.save(path)


def _pdf(path: Path, lines: list[str]) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    y = 82
    for index, line in enumerate(lines):
        page.insert_text((65, y), line, fontsize=17 if index == 0 else 12, fontname="china-s")
        y += 43 if index == 0 else 30
    document.save(path)


CASES = [
    {
        "id": "demo-qihang",
        "company": "启航供应链贸易有限公司",
        "amount": 500,
        "term": 12,
        "purpose": "采购电子元器件并补充订单周转资金",
        "materials": [
            ("demo-qihang-application", "授信申请与经营情况.docx", "融资申请", "docx", [
                "启航供应链贸易有限公司授信申请",
                "企业申请500万元融资担保，期限12个月，用于采购电子元器件并补充订单周转资金。",
                "公司主营工业电子元器件分销，第一大客户占2025年销售收入33%。",
            ]),
            ("demo-qihang-finance", "财务报表_横向年度.xlsx", "财务报表", "xlsx", {
                "年度数据": [["年度", "营业收入", "净利润", "经营活动现金流净额", "应收账款", "资产总额", "负债总额", "流动资产", "流动负债", "有息负债"],
                           [2024, 8000, 240, 80, 1100, 5200, 3500, 3400, 2600, 1800],
                           [2025, 9500, 280, -100, 1900, 6000, 4500, 4000, 3500, 2200]],
            }),
            ("demo-qihang-reconcile", "税票流水核验.xlsx", "纳税资料", "xlsx", {
                "核验": [["年度", "纳税申报收入", "银行经营性流入", "用途合同合计"], [2025, 9400, 7200, 550]],
                "债务": [["指标", "金额（万元）"], ["未来12个月到期债务", 1300], ["无合同关联方往来", 0]],
            }),
            ("demo-qihang-risk", "抵押及访谈摘要.docx", "抵押物", "docx", [
                "抵押及访谈摘要",
                "拟抵押仓储设备折后抵押价值600万元，评估报告出具于8个月前。",
                "已提供采购合同金额合计550万元，能够覆盖本次500万元申请。",
                "访谈确认未来12个月到期债务1300万元，现有有息负债2200万元。",
            ]),
        ],
        "sources": [
            ("s-q-app", "demo-qihang-application", "docx_paragraph", "第 3 段", "申请500万元融资担保，期限12个月，用于采购电子元器件并补充订单周转资金。"),
            ("s-q-customer", "demo-qihang-application", "docx_paragraph", "第 4 段", "2025年第一大客户销售占比33%。"),
            ("s-q-fin24", "demo-qihang-finance", "xlsx_range", "年度数据!A2:J2", "2024 | 8000 | 240 | 80 | 1100 | 5200 | 3500 | 3400 | 2600 | 1800"),
            ("s-q-fin25", "demo-qihang-finance", "xlsx_range", "年度数据!A3:J3", "2025 | 9500 | 280 | -100 | 1900 | 6000 | 4500 | 4000 | 3500 | 2200"),
            ("s-q-reconcile", "demo-qihang-reconcile", "xlsx_range", "核验!A2:D2", "2025 | 纳税申报收入9400万元 | 银行经营性流入7200万元 | 用途合同合计550万元"),
            ("s-q-debt", "demo-qihang-reconcile", "xlsx_range", "债务!A2:B3", "未来12个月到期债务1300万元；无合同关联方往来0万元。"),
            ("s-q-collateral", "demo-qihang-risk", "docx_paragraph", "第 3 段", "折后抵押价值600万元，评估报告出具于8个月前。"),
        ],
        "facts": [
            ("revenue", "2024", 8000, "万元", "s-q-fin24"), ("revenue", "2025", 9500, "万元", "s-q-fin25"),
            ("net_profit", "2025", 280, "万元", "s-q-fin25"), ("operating_cash_flow", "2025", -100, "万元", "s-q-fin25"),
            ("accounts_receivable", "2024", 1100, "万元", "s-q-fin24"), ("accounts_receivable", "2025", 1900, "万元", "s-q-fin25"),
            ("total_assets", "2025", 6000, "万元", "s-q-fin25"), ("total_liabilities", "2025", 4500, "万元", "s-q-fin25"),
            ("current_assets", "2025", 4000, "万元", "s-q-fin25"), ("current_liabilities", "2025", 3500, "万元", "s-q-fin25"),
            ("interest_bearing_debt", "2025", 2200, "万元", "s-q-fin25"), ("debt_due_12m", "2026", 1300, "万元", "s-q-debt"),
            ("tax_revenue", "2025", 9400, "万元", "s-q-reconcile"), ("operating_inflow", "2025", 7200, "万元", "s-q-reconcile"),
            ("top_customer_share", "2025", 33, "%", "s-q-customer"), ("related_unsupported_amount", "2025", 0, "万元", "s-q-debt"),
            ("supporting_contract_amount", "2026", 550, "万元", "s-q-reconcile"),
            ("collateral_discounted_value", "2026", 600, "万元", "s-q-collateral"),
            ("collateral_valuation_age_months", "2026", 8, "个月", "s-q-collateral"),
        ],
        "expected_warnings": ["revenue_bank", "ar_growth", "cash_profit", "leverage", "customer", "debt_due"],
    },
    {
        "id": "demo-xingqiao",
        "company": "星桥数字科技服务有限公司",
        "amount": 300,
        "term": 18,
        "purpose": "补充软件项目研发及人员成本",
        "materials": [
            ("demo-xingqiao-application", "担保申请.pdf", "融资申请", "pdf", [
                "星桥数字科技服务有限公司担保申请",
                "申请金额：300万元",
                "期限：18个月",
                "用途：补充软件项目研发及人员成本",
                "本材料为虚构演示数据。",
            ]),
            ("demo-xingqiao-company", "公司业务与客户说明.docx", "企业信息", "docx", [
                "星桥数字科技服务有限公司业务说明",
                "公司主营企业软件定制与运维服务。2025年第一大客户收入占比48%。",
                "向关联企业支付200万元往来款，未提供对应合同或服务成果。",
            ]),
            ("demo-xingqiao-data", "年度经营数据_纵向.xlsx", "财务报表", "xlsx", {
                "财务": [["指标（万元）", 2024, 2025], ["营业收入", 2600, 3300], ["净利润", 320, 400],
                       ["经营活动现金流净额", 210, 300], ["应收账款", 700, 850], ["资产总额", 2600, 3000],
                       ["负债总额", 1100, 1200], ["流动资产", 1700, 2000], ["流动负债", 750, 800], ["有息负债", 650, 700]],
                "外部口径": [["指标", "2025金额（万元）"], ["纳税申报收入", 2700], ["银行经营性流入", 3500],
                          ["用途合同合计", 180], ["未来12个月到期债务", 200]],
            }),
            ("demo-xingqiao-debt", "债务及用途核验.docx", "负债担保", "docx", [
                "债务及用途核验",
                "现有有息负债700万元，未来12个月到期债务200万元。",
                "已提供研发外包及人员服务合同180万元，低于300万元融资申请。",
                "企业未提供抵押物或其他反担保估值材料。",
            ]),
        ],
        "sources": [
            ("s-x-app", "demo-xingqiao-application", "pdf_page", "第 1 页", "申请金额300万元，期限18个月，用于补充软件项目研发及人员成本。"),
            ("s-x-company", "demo-xingqiao-company", "docx_paragraph", "第 3 段", "公司主营企业软件定制与运维服务，2025年第一大客户收入占比48%。"),
            ("s-x-related", "demo-xingqiao-company", "docx_paragraph", "第 4 段", "向关联企业支付200万元往来款，未提供对应合同。"),
            ("s-x-fin", "demo-xingqiao-data", "xlsx_range", "财务!A2:C10", "2024/2025营业收入2600/3300；净利润320/400；经营现金流210/300；应收700/850；资产3000；负债1200；流动资产2000；流动负债800；有息负债700。"),
            ("s-x-external", "demo-xingqiao-data", "xlsx_range", "外部口径!A2:B6", "2025纳税申报收入2700万元；银行经营性流入3500万元；用途合同180万元；未来12个月到期债务200万元。"),
            ("s-x-debt", "demo-xingqiao-debt", "docx_paragraph", "第 3-5 段", "有息负债700万元，到期债务200万元；用途合同180万元；未提供抵押物估值材料。"),
        ],
        "facts": [
            ("revenue", "2024", 2600, "万元", "s-x-fin"), ("revenue", "2025", 3300, "万元", "s-x-fin"),
            ("net_profit", "2025", 400, "万元", "s-x-fin"), ("operating_cash_flow", "2025", 300, "万元", "s-x-fin"),
            ("accounts_receivable", "2024", 700, "万元", "s-x-fin"), ("accounts_receivable", "2025", 850, "万元", "s-x-fin"),
            ("total_assets", "2025", 3000, "万元", "s-x-fin"), ("total_liabilities", "2025", 1200, "万元", "s-x-fin"),
            ("current_assets", "2025", 2000, "万元", "s-x-fin"), ("current_liabilities", "2025", 800, "万元", "s-x-fin"),
            ("interest_bearing_debt", "2025", 700, "万元", "s-x-fin"), ("debt_due_12m", "2026", 200, "万元", "s-x-external"),
            ("tax_revenue", "2025", 2700, "万元", "s-x-external"), ("operating_inflow", "2025", 3500, "万元", "s-x-external"),
            ("top_customer_share", "2025", 48, "%", "s-x-company"), ("related_unsupported_amount", "2025", 200, "万元", "s-x-related"),
            ("supporting_contract_amount", "2026", 180, "万元", "s-x-external"),
        ],
        "expected_warnings": ["revenue_tax", "customer", "related", "purpose"],
    },
]


def generate_portfolio_materials() -> list[dict]:
    root = settings.resolved_data_dir / "demo_materials"
    root.mkdir(parents=True, exist_ok=True)
    output = []
    for case in CASES:
        directory = root / case["id"]
        directory.mkdir(exist_ok=True)
        for document_id, filename, kind, file_type, content in case["materials"]:
            path = directory / filename
            if not path.exists():
                if file_type == "docx":
                    _docx(path, content[0], content[1:])
                elif file_type == "xlsx":
                    _xlsx(path, content)
                else:
                    _pdf(path, content)
            output.append({"id": document_id, "filename": filename, "kind": kind, "path": path})
    return output


def seed_portfolio_cases() -> None:
    material_map = {item["id"]: item for item in generate_portfolio_materials()}
    now = now_iso()
    with connect() as conn:
        for case in CASES:
            if conn.execute("SELECT 1 FROM projects WHERE id=?", (case["id"],)).fetchone():
                continue
            conn.execute(
                "INSERT INTO projects (id,access_token,company_name,amount,term_months,purpose,status,created_at,updated_at,consent_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (case["id"], secrets.token_urlsafe(24), case["company"], case["amount"], case["term"], case["purpose"], "ready", now, now, now),
            )
            for document_id, filename, kind, _, _ in case["materials"]:
                material = material_map[document_id]
                path: Path = material["path"]
                media = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}[path.suffix]
                locator_count = sum(source[1] == document_id for source in case["sources"])
                conn.execute(
                    "INSERT INTO documents (id,project_id,filename,kind,media_type,sha256,size,parse_status,quality,locator_count,warning,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    (document_id, case["id"], filename, kind, media, sha256(path), path.stat().st_size, "parsed", 0.96, locator_count, "", now),
                )
            for source_id, document_id, locator_type, locator, text in case["sources"]:
                conn.execute("INSERT INTO sources (id,project_id,document_id,locator_type,locator,text) VALUES (?,?,?,?,?,?)",
                             (source_id, case["id"], document_id, locator_type, locator, text))
            for index, (fact_key, period, value, unit, source_id) in enumerate(case["facts"], start=1):
                definition = FACT_REGISTRY[fact_key]
                conn.execute(
                    "INSERT INTO facts (id,project_id,category,name,period,value,unit,text_value,source_ids,confidence,fact_key,raw_value,extraction_method,review_status) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (f"f-{case['id']}-{index}", case["id"], definition.category, definition.label, period, value, unit, "",
                     json_dump([source_id]), 0.99, fact_key, str(value), "seed", "unreviewed"),
                )


def evaluation_manifest() -> dict[str, dict]:
    return {
        case["id"]: {
            "expected_fact_count": len(case["facts"]),
            "expected_facts": [
                {"fact_key": fact_key, "period": period, "value": value, "unit": unit}
                for fact_key, period, value, unit, _ in case["facts"]
            ],
            "expected_warning_keys": case["expected_warnings"],
        }
        for case in CASES
    }
