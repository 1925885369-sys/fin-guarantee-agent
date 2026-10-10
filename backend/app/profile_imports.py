from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import settings
from .database import connect, json_dump, json_load, now_iso, row, rows
from .entity_profiles import create_profile, get_profile, update_profile
from .facts import FACT_REGISTRY
from .parsers import parse_document
from .profile_catalog import RECORD_TYPES
from .schemas import EntityProfileCreate, EntityProfileUpdate, ProfileImportCommit


def _clean(value: Any) -> Any:
    if value is None:
        return ""
    return value.strip() if isinstance(value, str) else value


def _sheet_rows(sheet) -> list[dict[str, Any]]:
    values = [[_clean(cell.value) for cell in row_] for row_ in sheet.iter_rows()]
    values = [row_ for row_ in values if any(value != "" for value in row_)]
    if not values:
        return []
    headers = [str(value).strip() for value in values[0]]
    output = []
    for values_row in values[1:]:
        item = {header: values_row[index] if index < len(values_row) else "" for index, header in enumerate(headers) if header}
        if any(value != "" for value in item.values()):
            output.append(item)
    return output


def _parse_xlsx(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    payload: dict[str, Any] = {"facts": [], "sources": [], **{key: [] for key in RECORD_TYPES}}
    for sheet in workbook.worksheets:
        key = sheet.title.strip().lower()
        if key in {"profile", "主体", "主体信息"}:
            rows_ = [[_clean(cell.value) for cell in row_] for row_ in sheet.iter_rows()]
            rows_ = [row_ for row_ in rows_ if any(value != "" for value in row_)]
            if rows_ and len(rows_[0]) == 2 and str(rows_[0][0]).lower() not in {"name", "名称"}:
                payload.update({str(item[0]).strip(): item[1] for item in rows_ if len(item) >= 2 and item[0]})
            elif rows_:
                headers = [str(value).strip() for value in rows_[0]]
                if len(rows_) > 1:
                    payload.update({header: rows_[1][index] for index, header in enumerate(headers) if header and index < len(rows_[1])})
        elif key in {"facts", "事实"}:
            payload["facts"].extend(_sheet_rows(sheet))
        elif key in {"sources", "来源"}:
            payload["sources"].extend(_sheet_rows(sheet))
        elif key in RECORD_TYPES:
            payload[key].extend({"data": item} for item in _sheet_rows(sheet))
    aliases = {
        "企业名称": "name", "名称": "name", "统一社会信用代码": "unified_code",
        "资料基准日": "baseline_date", "最新财务期间": "latest_financial_period",
    }
    for source, target in aliases.items():
        if source in payload and target not in payload:
            payload[target] = payload.pop(source)
    return payload


def _parse_structured(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        try:
            payload = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("JSON 文件无法解析") from exc
        if not isinstance(payload, dict):
            raise ValueError("JSON 顶层必须是对象")
        return payload
    if suffix == ".xlsx":
        try:
            return _parse_xlsx(content)
        except Exception as exc:
            raise ValueError("Excel 文件无法解析") from exc
    raise ValueError("结构化导入仅支持 JSON 或 XLSX")


def _normalize_source(item: dict[str, Any], index: int) -> dict[str, Any]:
    source_id = str(item.get("id") or f"import-source-{index}")
    page = item.get("page")
    if page in {"", None}:
        page = None
    return {
        "id": source_id, "filename": str(item.get("filename") or "结构化导入"),
        "kind": str(item.get("kind") or "主体档案"), "locator": str(item.get("locator") or ""),
        "locator_type": str(item.get("locator_type") or "profile_locator"),
        "page": int(page) if page is not None else None, "quote": str(item.get("quote") or ""),
        "content_hash": str(item.get("content_hash") or ""),
    }


def _normalize_fact(item: dict[str, Any]) -> dict[str, Any]:
    source_ids = item.get("source_ids") or []
    if isinstance(source_ids, str):
        source_ids = [value.strip() for value in re.split(r"[,，;；]", source_ids) if value.strip()]
    value = item.get("value")
    if value == "":
        value = None
    return {
        "fact_key": str(item.get("fact_key") or item.get("field") or ""),
        "category": str(item.get("category") or ""), "value": float(value) if value is not None else None,
        "unit": str(item.get("unit") or ""), "text_value": str(item.get("text_value") or ""),
        "period": str(item.get("period") or ""), "as_of_date": item.get("as_of_date") or None,
        "period_type": str(item.get("period_type") or "unknown"), "scope": str(item.get("scope") or "entity"),
        "subject_role": str(item.get("subject_role") or "issuer"),
        "status": str(item.get("status") or "single_source"), "valid_until": item.get("valid_until") or None,
        "source_ids": source_ids, "raw_value": str(item.get("raw_value") or ""),
        "confidence": float(item.get("confidence") or 0.8), "currency": str(item.get("currency") or "CNY"),
        "reviewer": str(item.get("reviewer") or ""), "reviewed_at": item.get("reviewed_at") or None,
    }


def _candidate(candidate_type: str, data: dict[str, Any], index: int, record_type: str = "") -> dict[str, Any]:
    digest = hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest()[:16]
    return {"id": f"candidate-{candidate_type}-{index}-{digest}", "type": candidate_type, "record_type": record_type, "data": data}


def create_import_job(
    filename: str, content: bytes, *, profile_id: str | None = None,
    evidence_files: list[tuple[str, str, bytes]] | None = None,
    package_kind: str = "", baseline_date: str | None = None,
    period: str | None = None, scope: str = "entity",
) -> dict[str, Any]:
    if profile_id and not get_profile(profile_id):
        raise LookupError("主体档案不存在")
    payload = _parse_structured(filename, content)
    candidates: list[dict[str, Any]] = []
    issues: list[dict[str, str]] = []
    metadata = {}
    for key in ("name", "unified_code", "baseline_date", "latest_financial_period", "data"):
        value = payload.get(key)
        if value is not None and value != "":
            metadata[key] = value
    if metadata:
        candidates.append(_candidate("profile", metadata, 1))
    for index, source in enumerate(payload.get("sources") or [], start=1):
        normalized = _normalize_source(source, index)
        candidates.append(_candidate("source", normalized, index))
    for index, fact in enumerate(payload.get("facts") or [], start=1):
        try:
            normalized = _normalize_fact(fact)
        except (TypeError, ValueError):
            issues.append({"key": f"fact:{index}", "message": "事实数值或字段格式无效"})
            continue
        if normalized["fact_key"] not in FACT_REGISTRY:
            issues.append({"key": f"fact:{index}", "message": f"未注册事实键：{normalized['fact_key']}"})
        candidates.append(_candidate("fact", normalized, index))
    for record_type in RECORD_TYPES:
        for index, record in enumerate(payload.get(record_type) or [], start=1):
            data = record if isinstance(record, dict) else {"data": {"value": record}}
            candidates.append(_candidate("record", data, index, record_type))

    job_id = str(uuid.uuid4())
    evidence_meta = []
    directory = settings.resolved_data_dir / "profile_imports" / job_id
    for index, (evidence_name, media_type, evidence_content) in enumerate(evidence_files or [], start=1):
        safe_name = re.sub(r"[^\w\u4e00-\u9fff.\-]+", "_", Path(evidence_name).name)[:120] or f"evidence-{index}"
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{index}-{safe_name}"
        path.write_bytes(evidence_content)
        digest = hashlib.sha256(evidence_content).hexdigest()
        file_id = str(uuid.uuid4())
        evidence_meta.append({"id": file_id, "filename": evidence_name, "media_type": media_type, "sha256": digest})
        parse_status = "stored"
        try:
            parsed = parse_document(path)
            parse_status = "parsed"
            for fragment_index, fragment in enumerate(parsed.fragments, start=1):
                source_data = {
                    "id": f"evidence-{digest[:12]}-{fragment_index}", "filename": evidence_name,
                    "kind": "主体档案", "locator": fragment.locator, "locator_type": fragment.locator_type,
                    "page": fragment.page, "quote": fragment.text[:4000],
                    "content_hash": hashlib.sha256(fragment.text.encode("utf-8")).hexdigest(),
                }
                candidates.append(_candidate("source", source_data, len(candidates) + 1))
        except Exception as exc:
            parse_status = "failed"
            issues.append({"key": f"evidence:{index}", "message": f"证据文件解析失败：{str(exc)[:120]}"})
        with connect() as conn:
            conn.execute(
                "INSERT INTO entity_profile_import_files (id,import_id,filename,media_type,sha256,size,storage_path,file_role,parse_status,created_at) "
                "VALUES (?,?,?,?,?,?,?,?,?,?)",
                (file_id, job_id, evidence_name, media_type, digest, len(evidence_content), str(path), "evidence", parse_status, now_iso()),
            )
    preview = {
        "candidates": candidates,
        "issues": issues,
        "evidence_files": evidence_meta,
        "context": {
            "package_kind": package_kind,
            "baseline_date": baseline_date,
            "period": period,
            "scope": scope,
        },
    }
    digest = hashlib.sha256(content).hexdigest()
    with connect() as conn:
        conn.execute(
            "INSERT INTO entity_profile_imports (id,profile_id,filename,import_format,sha256,status,payload_json,preview_json,error,created_at) "
            "VALUES (?,?,?,?,?,'preview',?,?,?,?)",
            (job_id, profile_id, filename, Path(filename).suffix.lower().lstrip("."), digest,
             json_dump(payload), json_dump(preview), "", now_iso()),
        )
    return get_import_job(job_id) or {}


def get_import_job(job_id: str) -> dict[str, Any] | None:
    item = row("SELECT * FROM entity_profile_imports WHERE id=?", (job_id,))
    if not item:
        return None
    item["payload"] = json_load(item.pop("payload_json"))
    item["preview"] = json_load(item.pop("preview_json"))
    storage_status = item["status"]
    preview = item["preview"]
    source_candidates = {
        candidate["data"].get("id"): candidate["data"]
        for candidate in preview.get("candidates", [])
        if candidate.get("type") == "source" and candidate.get("data", {}).get("id")
    }
    frontend_candidates = []
    for candidate in preview.get("candidates", []):
        if candidate.get("type") != "fact":
            continue
        data = candidate.get("data", {})
        definition = FACT_REGISTRY.get(data.get("fact_key", ""))
        source_ids = [str(value) for value in data.get("source_ids") or []]
        primary = next((source_candidates[value] for value in source_ids if value in source_candidates), {})
        issue_codes = []
        if not source_ids or not primary or not str(primary.get("locator") or "").strip() or not str(primary.get("quote") or "").strip():
            issue_codes.append("missing_evidence")
        if not definition:
            issue_codes.append("unregistered_fact_key")
        frontend_candidates.append({
            "id": candidate["id"],
            "fact_key": data.get("fact_key", ""),
            "name": definition.label if definition else data.get("fact_key", ""),
            "category": data.get("category") or (definition.category if definition else ""),
            "value": data.get("value"),
            "text_value": data.get("text_value", ""),
            "unit": data.get("unit", ""),
            "period": data.get("period") or None,
            "scope": data.get("scope", "entity"),
            "proposed_status": data.get("status", "single_source"),
            "source_ids": source_ids,
            "source_label": primary.get("filename", ""),
            "locator": primary.get("locator", ""),
            "quote": primary.get("quote", ""),
            "confidence": float(data.get("confidence") or 0),
            "issue_codes": issue_codes,
        })
    files = [{
        "id": f"{job_id}-structured",
        "filename": item["filename"],
        "size": 0,
        "status": "parsed",
        "page_count": None,
        "parse_quality": 1.0,
        "error": "",
    }]
    import_files = rows(
        "SELECT * FROM entity_profile_import_files WHERE import_id=? ORDER BY created_at,id", (job_id,),
    )
    for file_row in import_files:
        parsed = file_row.get("parse_status") == "parsed"
        files.append({
            "id": file_row["id"], "filename": file_row["filename"], "size": int(file_row.get("size") or 0),
            "status": "parsed" if parsed else "failed" if file_row.get("parse_status") == "failed" else "queued",
            "page_count": None, "parse_quality": 1.0 if parsed else None,
            "error": "" if parsed else ("证据文件解析失败" if file_row.get("parse_status") == "failed" else ""),
        })
    frontend_issues = [{
        "code": issue.get("key", "import_issue"),
        "severity": "error" if "失败" in str(issue.get("message", "")) else "warning",
        "message": issue.get("message", ""),
    } for issue in preview.get("issues", [])]
    context = preview.get("context", {})
    mapped_status = storage_status
    if storage_status == "preview":
        mapped_status = "ready_to_commit" if frontend_candidates and not frontend_issues else "needs_mapping"
    item.update({
        "storage_status": storage_status,
        "status": mapped_status,
        "target_profile_id": item.get("profile_id"),
        "package_kind": context.get("package_kind", ""),
        "baseline_date": context.get("baseline_date"),
        "period": context.get("period"),
        "scope": context.get("scope", "entity"),
        "progress": 100 if storage_status in {"preview", "committed"} else 0,
        "files": files,
        "candidates": frontend_candidates,
        "issues": frontend_issues,
    })
    return item


def _merge_by_key(existing: list[dict[str, Any]], incoming: list[dict[str, Any]], key_fields: tuple[str, ...]) -> list[dict[str, Any]]:
    output = {tuple(str(item.get(field) or "") for field in key_fields): item for item in existing}
    for item in incoming:
        output[tuple(str(item.get(field) or "") for field in key_fields)] = item
    return list(output.values())


def commit_import_job(job_id: str, selection: ProfileImportCommit) -> dict[str, Any]:
    job = get_import_job(job_id)
    if not job:
        raise LookupError("导入任务不存在")
    if job.get("storage_status", job["status"]) != "preview":
        raise ValueError("导入任务已经提交或失效")
    candidates = {item["id"]: item for item in job["preview"].get("candidates", [])}
    selected_ids = set(selection.selected_candidate_ids)
    unknown = sorted(selected_ids - set(candidates))
    if unknown:
        raise ValueError(f"导入候选不存在：{','.join(unknown)}")
    selected = [candidates[item_id] for item_id in selection.selected_candidate_ids]
    referenced_source_ids = {
        source_id
        for item in selected if item.get("type") in {"fact", "record"}
        for source_id in item.get("data", {}).get("source_ids", [])
    }
    selected_ids = {item["id"] for item in selected}
    selected.extend(
        item for item in candidates.values()
        if item.get("type") == "source"
        and item.get("data", {}).get("id") in referenced_source_ids
        and item["id"] not in selected_ids
    )
    for correction in selection.corrections:
        if correction.candidate_id not in candidates:
            raise ValueError(f"修正候选不存在：{correction.candidate_id}")
        original = candidates[correction.candidate_id]
        selected = [item for item in selected if item["id"] != correction.candidate_id]
        selected.append({**original, "data": {**original.get("data", {}), **correction.patch_data()}})

    profile_id = selection.profile_id or job.get("profile_id")
    current = get_profile(profile_id) if profile_id else None
    if profile_id and not current:
        raise LookupError("主体档案不存在")
    metadata = next((item["data"] for item in selected if item["type"] == "profile"), {})
    sources = [item["data"] for item in selected if item["type"] == "source"]
    facts = [item["data"] for item in selected if item["type"] == "fact"]
    records = {
        record_type: [item["data"] for item in selected if item["type"] == "record" and item["record_type"] == record_type]
        for record_type in RECORD_TYPES
    }
    if current:
        sources = _merge_by_key(current["sources"], sources, ("id",))
        facts = _merge_by_key(current["facts"], facts, ("fact_key", "period", "scope"))
        for record_type in RECORD_TYPES:
            records[record_type] = [*current.get(record_type, []), *records[record_type]]
    known_source_ids = {item.get("id") for item in sources if item.get("id")}
    for item in [*facts, *(record for record_type in RECORD_TYPES for record in records[record_type])]:
        missing = set(item.get("source_ids") or []) - known_source_ids
        if missing:
            raise ValueError(f"提交记录引用未选中的来源：{','.join(sorted(missing))}")
    payload = {
        "name": metadata.get("name") or (current or {}).get("name"),
        "unified_code": metadata.get("unified_code") or (current or {}).get("unified_code"),
        "baseline_date": metadata.get("baseline_date") or (current or {}).get("baseline_date"),
        "latest_financial_period": metadata.get("latest_financial_period") or (current or {}).get("latest_financial_period"),
        "data": metadata.get("data") or (current or {}).get("data", {}),
        "sources": sources, "facts": facts, "change_note": selection.change_note, **records,
    }
    if current:
        validated = EntityProfileUpdate.model_validate(payload)
        result = update_profile(current["id"], validated.model_dump(exclude_none=True))
    else:
        validated = EntityProfileCreate.model_validate(payload)
        result = create_profile(validated.model_dump())
    with connect() as conn:
        conn.execute(
            "UPDATE entity_profile_imports SET profile_id=?,status='committed',committed_version=?,committed_at=? WHERE id=?",
            (result["id"], result["version"], now_iso(), job_id),
        )
    return {
        **result,
        "job_id": job_id,
        "profile_id": result["id"],
        "import_status": "committed",
    }
