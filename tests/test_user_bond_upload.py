from __future__ import annotations

from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from backend.app.config import settings
from backend.app.database import connect, row
from backend.app.main import app
from backend.app.workflow import _node_input_hash


def _write_docx(path: Path, heading: str, paragraphs: list[str]) -> None:
    document = Document()
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def _upload(client: TestClient, project_id: str, token: str, path: Path):
    with path.open("rb") as stream:
        return client.post(
            f"/api/projects/{project_id}/documents",
            headers={"X-Project-Token": token},
            files={"file": (path.name, stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )


def _delete_project(project_id: str) -> None:
    tables = (
        "llm_attempts", "run_artifacts", "audit_events", "evidence_edges", "report_claims",
        "gold_annotations", "candidate_findings", "review_tasks", "policy_snapshots",
        "financial_tie_outs", "financial_metrics", "project_entities", "canonical_facts",
        "fact_assertions", "bond_tranches", "run_nodes", "runs", "report_sections", "risks",
        "cross_checks", "facts", "citations", "sources", "documents",
    )
    with connect() as conn:
        for table in tables:
            conn.execute(f"DELETE FROM {table} WHERE project_id=?", (project_id,))
        conn.execute("DELETE FROM projects WHERE id=?", (project_id,))
    upload_dir = settings.resolved_data_dir / "uploads" / project_id
    if upload_dir.is_dir():
        for path in upload_dir.iterdir():
            if path.is_file():
                path.unlink()
        upload_dir.rmdir()


def test_user_bond_core_material_upload_classification_and_duplicate_control(tmp_path: Path):
    prospectus = tmp_path / "测试公司募集说明书.docx"
    rating = tmp_path / "测试公司信用评级报告.docx"
    _write_docx(prospectus, "募集说明书", ["发行条款", "本期债券发行规模为2亿元，期限3年。"])
    _write_docx(rating, "信用评级报告", ["评级结果", "主体评级AA-，债项评级AA+，评级展望稳定。"])

    project_id = ""
    with TestClient(app) as client:
        created = client.post("/api/projects", json={
            "company_name": "普通上传测试有限公司",
            "amount": 20000,
            "term_months": 36,
            "purpose": "债券发行融资担保完整尽调",
            "data_consent": True,
            "project_type": "bond_guarantee",
            "bond_name": "普通上传测试有限公司2026年公司债券",
        })
        assert created.status_code == 201
        project_id = created.json()["id"]
        token = created.json()["accessToken"]
        headers = {"X-Project-Token": token}
        try:
            uploaded_prospectus = _upload(client, project_id, token, prospectus)
            uploaded_rating = _upload(client, project_id, token, rating)
            assert uploaded_prospectus.status_code == 201
            assert uploaded_rating.status_code == 201
            assert uploaded_prospectus.json()["documentRole"] == "prospectus"
            assert uploaded_rating.json()["documentRole"] == "rating_report"

            inventory = client.get(f"/api/projects/{project_id}/document-inventory", headers=headers).json()
            assert inventory["complete_due_diligence_ready"] is True
            assert inventory["analysis_mode"] == "full_due_diligence"
            assert {item["document_role"] for item in inventory["documents"]} == {"prospectus", "rating_report"}

            duplicate = _upload(client, project_id, token, prospectus)
            assert duplicate.status_code == 409
            assert duplicate.json()["detail"]["code"] == "duplicate_document"
            assert duplicate.json()["detail"]["document_id"] == uploaded_prospectus.json()["id"]
            assert row("SELECT COUNT(*) count FROM documents WHERE project_id=?", (project_id,))["count"] == 2

            with connect() as conn:
                conn.execute(
                    "INSERT INTO facts (id,project_id,category,name,period,value,unit,text_value,source_ids,confidence,fact_key,raw_value,extraction_method,review_status,subject,subject_id,scope,period_type) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    ("user-upload-system-fact", project_id, "债券基本信息", "债券名称", "", None, "", "旧系统值", "[]", 0.8, "bond_name", "旧系统值", "model", "unreviewed", "", "", "transaction", "undated"),
                )
                conn.execute(
                    "INSERT INTO facts (id,project_id,category,name,period,value,unit,text_value,source_ids,confidence,fact_key,raw_value,extraction_method,review_status,subject,subject_id,scope,period_type) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    ("user-upload-human-fact", project_id, "债券基本信息", "债券名称", "", None, "", "人工确认值", "[]", 1.0, "bond_name", "人工确认值", "model", "human_verified", "", "", "transaction", "undated"),
                )
                conn.execute(
                    "INSERT INTO report_sections (id,project_id,section_key,title,content,citation_ids,review_status,updated_at) VALUES (?,?,?,?,?,?,?,?)",
                    ("user-upload-old-report", project_id, "scope", "旧报告", "旧派生内容", "[]", "draft", "2026-01-01T00:00:00Z"),
                )
                conn.execute(
                    "INSERT INTO run_artifacts (id,run_id,project_id,node_key,input_hash,artifact_type,artifact_ids,status,source_run_id,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    ("user-upload-old-artifact", "old-run", project_id, "extract", "old-hash", "extract", "[]", "completed", "", "2026-01-01T00:00:00Z"),
                )
            before_hash = _node_input_hash(project_id, "extract")
            changed = client.patch(
                f"/api/projects/{project_id}/documents/{uploaded_prospectus.json()['id']}",
                headers=headers,
                json={"kind": "其他材料", "document_role": "other"},
            )
            assert changed.status_code == 200 and changed.json()["changed"] is True
            assert _node_input_hash(project_id, "extract") != before_hash
            inventory = client.get(f"/api/projects/{project_id}/document-inventory", headers=headers).json()
            assert inventory["complete_due_diligence_ready"] is False
            assert inventory["analysis_mode"] == "insufficient"
            assert row("SELECT id FROM facts WHERE id='user-upload-human-fact'")
            assert row("SELECT id FROM facts WHERE id='user-upload-system-fact'") is None
            assert row("SELECT id FROM report_sections WHERE id='user-upload-old-report'") is None
            assert row("SELECT status FROM run_artifacts WHERE id='user-upload-old-artifact'")["status"] == "stale"
            assert row(
                "SELECT id FROM audit_events WHERE project_id=? AND action='classify' AND entity_id=?",
                (project_id, uploaded_prospectus.json()["id"]),
            )

            mismatch = client.patch(
                f"/api/projects/{project_id}/documents/{uploaded_prospectus.json()['id']}",
                headers=headers,
                json={"kind": "募集说明书", "document_role": "rating_report"},
            )
            assert mismatch.status_code == 422
        finally:
            _delete_project(project_id)


def test_demo_document_classification_is_read_only():
    with TestClient(app) as client:
        inventory = client.get("/api/projects/demo-bond-huachen/document-inventory").json()
        document_id = inventory["documents"][0]["id"]
        response = client.patch(
            f"/api/projects/demo-bond-huachen/documents/{document_id}",
            json={"kind": "其他材料", "document_role": "other"},
        )
        assert response.status_code == 403
