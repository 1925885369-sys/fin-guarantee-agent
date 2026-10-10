"""报告章节 blocks 与 Markdown/Word 三出口一致性（Track A）。

守住四件事：

1. 章节 ``blocks`` 里的表格块**直接来自 ``TableBlock``**——把正文里的数值换掉，
   表格块仍然给已校验事实的数，证明它不是从 Markdown 反解出来的；
2. 段落块里**不含**系统表格块与冲突块（表名行、表体、表注、校验行、披露值列表），
   只留分析文字；
3. Markdown 与 Word 出口都保留表名、列、单位与表注，Word 表头跨页重复；
4. 导出件里没有禁用材料名、内部 UUID 与内部字段键（``wording_violations`` 为空）。
"""
from __future__ import annotations

import io
import json
import re

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend.app.database import connect, now_iso
from backend.app.delivery_wording import wording_violations
from backend.app.exports import docx_report, markdown_report
from backend.app.main import app
from backend.app.professional import sync_fact_assertions
from backend.app.report_blocks import TABLE_CAPTION_RE
from backend.app.report_delivery import write_report_section
from backend.app.report_tables import SYSTEM_TABLE_HEADING, chapter_table_blocks, system_table_markdown

SUBJECT = "表格轨道测试有限公司"
ORIGINAL_FILENAME = "募集说明书（申报稿）.pdf"
CONFLICT_BLOCK = (
    "### 冲突事实（以下披露值必须并列保留）\n"
    "- 营业收入（2025）：\n"
    "  - 披露值 30,000万元；来源：发行文件〔资料01〕第 3 页\n"
    "  - 披露值 33,000万元；来源：外部信用评价文件〔资料02〕第 5 页"
)

# (字段, 科目名, 2025 值)。数值故意取整数，便于在正文里做替换实验。
FACTS = (
    ("total_assets", "资产总计", 72000),
    ("total_liabilities", "负债合计", 50400),
    ("equity", "所有者权益合计", 21600),
    ("revenue", "一、营业收入", 30000),
    ("net_profit", "四、净利润", 3000),
)


@pytest.fixture(scope="module")
def project():
    """自建项目：不走演示项目，避免与其它用例共享报告章节。"""
    with TestClient(app) as client:
        created = client.post("/api/projects", json={
            "company_name": SUBJECT,
            "amount": 20000,
            "term_months": 36,
            "purpose": "债券发行融资担保完整尽调",
            "data_consent": True,
            "project_type": "bond_guarantee",
            "bond_name": f"{SUBJECT}2026年公司债券",
        })
        assert created.status_code == 201, created.text
        project_id, token = created.json()["id"], created.json()["accessToken"]
        document_id, source_id, citation_id = (f"doc-{project_id}", f"src-{project_id}", f"cit-{project_id}")
        now = now_iso()
        with connect() as conn:
            # 材料原件的文件名是工作台口径（禁用表述在这里是合法的），成品里必须换掉。
            conn.execute(
                "INSERT INTO documents (id,project_id,filename,kind,media_type,sha256,size,parse_status,"
                "quality,locator_count,warning,created_at,parser,ocr_status,page_count,document_role) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (document_id, project_id, ORIGINAL_FILENAME, "募集说明书", "application/pdf", "x", 1,
                 "parsed", 1.0, 1, "", now, "native", "skipped", 1, "prospectus"),
            )
            conn.execute(
                "INSERT INTO sources (id,project_id,document_id,locator_type,locator,text,page,bbox_json,"
                "extraction_method,content_hash) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (source_id, project_id, document_id, "page", "第 3 页",
                 f"{SUBJECT} 虚构案例，仅用于系统演示", 3, "{}", "native", "h"),
            )
            conn.execute(
                "INSERT INTO citations (id,project_id,source_id,document_id,label,quote,locator_type,locator) "
                "VALUES (?,?,?,?,?,?,?,?)",
                (citation_id, project_id, source_id, document_id, ORIGINAL_FILENAME,
                 f"{SUBJECT} 虚构案例，仅用于系统演示", "pdf_page", "第 3 页"),
            )
            for key, name, value in FACTS:
                for period, factor in (("2025", 1.0), ("2024", 0.9), ("2023", 0.8)):
                    conn.execute(
                        "INSERT INTO facts (id,project_id,category,name,period,value,unit,text_value,"
                        "source_ids,confidence,fact_key,raw_value,extraction_method,review_status,subject,"
                        "subject_id,scope,period_type) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (f"fact-{key}-{period}", project_id, "财务数据", name, period, value * factor,
                         "万元", "", json.dumps([source_id]), 0.9, key, str(value * factor), "model",
                         "unreviewed", "", "", "consolidated", "annual"),
                    )
        sync_fact_assertions(project_id)
        payload = {"id": project_id, "token": token, "client": client, "citation_id": citation_id}
        try:
            write_financial_section(payload)
            yield payload
        finally:
            with connect() as conn:
                for table in ("report_sections", "citations", "fact_assertions", "canonical_facts",
                              "facts", "table_cell_reviews", "review_tasks", "audit_events",
                              "sources", "documents"):
                    conn.execute(f"DELETE FROM {table} WHERE project_id=?", (project_id,))
                conn.execute("DELETE FROM project_entities WHERE project_id=?", (project_id,))
                conn.execute("DELETE FROM projects WHERE id=?", (project_id,))


def analysis_content(project_id: str, *, with_tables: bool = True) -> str:
    """章节正文：分析文字 + 系统表格块 + 冲突块（与真实生成时的拼装一致）。"""
    parts = ["### 本章判断", "", "资产负债规模随经营扩张同步增长，结构未见异常。", ""]
    if with_tables:
        parts.extend([system_table_markdown(project_id, "financial", subject=SUBJECT), ""])
    parts.extend(["### 表后分析", "", "负债以短期借款为主，需取得原件后核验期限结构。", "", CONFLICT_BLOCK])
    return "\n".join(parts)


def write_financial_section(project: dict, *, with_tables: bool = True) -> None:
    system_block = system_table_markdown(project["id"], "financial", subject=SUBJECT)
    write_report_section(
        project_id=project["id"], section_key="financial", title="财务分析", section_id=f"{project['id']}-financial",
        content=analysis_content(project["id"], with_tables=with_tables),
        citation_ids=[project["citation_id"]],
        content_origin="deterministic", generation_status="deterministic",
        required_disclosures=[block for block in (CONFLICT_BLOCK, system_block if with_tables else "") if block],
        deterministic_draft=analysis_content(project["id"], with_tables=with_tables),
        length_policy={"effective_min_chars": 1, "effective_max_chars": 20000},
        support_status="pending",
    )


def api_blocks(project: dict, section_key: str = "financial") -> list[dict]:
    response = project["client"].get(
        f"/api/projects/{project['id']}/report", headers={"X-Project-Token": project["token"]},
    )
    assert response.status_code == 200, response.text
    section = next(item for item in response.json() if item["section_key"] == section_key)
    # 原有字段必须原样保留：blocks 是新增的呈现层，不能顶掉正文与状态。
    for key in ("content", "content_origin", "delivery_status", "disclosure_status", "citation_ids"):
        assert key in section, key
    return section["blocks"]


def table_blocks(blocks: list[dict]) -> list[dict]:
    return [block for block in blocks if block["kind"] == "table"]


def docx_paragraphs(document: Document) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


def test_api_blocks_include_table_blocks_matching_table_block_data(project) -> None:
    blocks = api_blocks(project)
    expected = chapter_table_blocks(project["id"], "financial", subject=SUBJECT)
    assert expected, "前提：财务章应有系统表格"
    assert all(block.validation_status == "passed" for block in expected), [
        (block.title, block.validation_summary) for block in expected
    ]

    produced = table_blocks(blocks)
    assert [block["table_id"] for block in produced] == [block.table_id for block in expected]
    for api_block, source in zip(produced, expected):
        assert api_block["table"] == source.to_dict()
        assert api_block["validation_status"] == source.validation_status
        assert api_block["origin_mode"] == source.origin_mode
        assert api_block["table_type"] == source.table_type
        assert api_block["title"] == source.title
        assert api_block["unit"] == source.unit
        assert api_block["columns"] == [
            {"key": column.key, "label": column.label, "role": column.role, "unit": column.unit}
            for column in source.columns
        ]
        assert api_block["notes"] == list(source.notes)
        assert api_block["present_in_content"] is True
        if source.origin_mode == "uploaded_table" and source.title != "报表及明细资料覆盖情况":
            assert api_block["table_key"], "材料原表应给出可回查的业务键"


def test_refresh_tables_preserves_prose_and_invalidates_final_review(project, monkeypatch) -> None:
    """A manual table refresh updates only disclosures and never invokes the model."""
    from backend.app import llm

    monkeypatch.setattr(llm, "rewrite_section", lambda *args, **kwargs: pytest.fail("refresh must not call the model"))
    write_financial_section(project)
    with connect() as conn:
        conn.execute("DELETE FROM report_table_snapshots WHERE project_id=? AND chapter_key='financial'", (project["id"],))
        conn.execute(
            "UPDATE projects SET report_revision=4,review_revision=4,report_reviewed_revision=4 WHERE id=?",
            (project["id"],),
        )
    before_response = project["client"].get(
        f"/api/projects/{project['id']}/report", headers={"X-Project-Token": project["token"]},
    )
    before = next(item for item in before_response.json() if item["section_key"] == "financial")
    response = project["client"].post(
        f"/api/projects/{project['id']}/report/tables/refresh",
        headers={"X-Project-Token": project["token"]},
        json={
            "chapters": ["financial"], "expected_report_revision": 4,
            "section_hashes": {"financial": before["content_sha256"]},
        },
    )
    assert response.status_code == 200, response.text
    after = project["client"].get(
        f"/api/projects/{project['id']}/report", headers={"X-Project-Token": project["token"]},
    )
    section = next(item for item in after.json() if item["section_key"] == "financial")
    assert "资产负债规模随经营扩张同步增长" in section["content"]
    assert section["content_origin"] == "deterministic"
    assert section["review_status"] == "draft"
    assert section["report_revision"] == 5
    assert response.json()["report_reviewed_revision"] == 0


def test_real_reference_table_refresh_is_local_only(project, monkeypatch) -> None:
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    import backend.app.main as main

    monkeypatch.setattr(main, "settings", SimpleNamespace(public_demo=False))
    payload = {
        "chapters": ["financial"], "expected_report_revision": 0,
        "section_hashes": {"financial": "section-hash"},
    }
    external = project["client"].post(
        "/api/projects/real-bond-zhongyou-2026-02/report/tables/refresh", json=payload,
    )
    assert external.status_code == 403
    with TestClient(app, client=("127.0.0.1", 50000)) as local_client:
        local = local_client.post(
            "/api/projects/real-bond-zhongyou-2026-02/report/tables/refresh", json=payload,
        )
    assert local.status_code != 403


def test_table_block_values_are_not_back_parsed_from_markdown(project) -> None:
    """把正文里的数值改掉：表格块仍给已校验事实的数，且标出正文与数据不一致。"""
    write_financial_section(project)
    source = chapter_table_blocks(project["id"], "financial", subject=SUBJECT)[0]
    cell = next(row for row in source.rows if row.row_key == "资产总计").cell(1)
    assert cell.raw == "72,000" and cell.normalized == "72000"

    blocks = table_blocks(api_blocks(project))
    first = blocks[0]
    # Markdown 反解拿不到这三样东西：规范化值、逐格事实字段、逐格来源。
    assert first["table"]["rows"][0]["cells"][1]["normalized"] == cell.normalized
    assert first["table"]["rows"][0]["cells"][1]["fact_key"] == "total_assets"
    assert first["table"]["rows"][0]["cells"][1]["source_ids"] == list(cell.source_ids)
    assert "|" not in json.dumps(first["table"]["rows"][0]["cells"][1]["normalized"])

    with connect() as conn:
        conn.execute(
            "UPDATE report_sections SET content=REPLACE(content,'72,000','888') WHERE project_id=? AND section_key='financial'",
            (project["id"],),
        )
    try:
        mutated = table_blocks(api_blocks(project))[0]
        # 正文里已经没有这张表了（数值被替换），但表格块的数据仍来自 TableBlock。
        assert mutated["table"]["rows"][0]["cells"][1]["raw"] == "72,000"
        assert mutated["table"] == first["table"]
        assert mutated["present_in_content"] is False
    finally:
        write_financial_section(project)


def test_paragraph_blocks_exclude_system_table_and_conflict_blocks(project) -> None:
    write_financial_section(project)
    blocks = api_blocks(project)
    paragraphs = [block.get("text") or "" for block in blocks if block["kind"] == "paragraph"]
    others = [block for block in blocks if block["kind"] != "table"]

    # 分析文字保留（表前与表后都在），且按原顺序排列。
    assert any("资产负债规模随经营扩张同步增长" in text for text in paragraphs)
    assert any("负债以短期借款为主" in text for text in paragraphs)
    assert paragraphs.index(next(t for t in paragraphs if "资产负债规模随经营扩张同步增长" in t)) < \
        paragraphs.index(next(t for t in paragraphs if "负债以短期借款为主" in t))

    for block in others:
        texts = [block.get("text") or "", block.get("label") or "", *(block.get("items") or [])]
        joined = "\n".join(texts).strip()
        assert SYSTEM_TABLE_HEADING not in joined
        assert not joined.startswith("系统整理数据表")
        assert not TABLE_CAPTION_RE.match(joined)
        # 系统表格块的 Markdown 行（表体）不得作为段落出现。
        assert "| --- |" not in joined and not joined.startswith("|")
        # 表注与冲突披露值都不得当成段落。
        assert "不是从原始报表逐格提取" not in joined
        assert "披露值" not in joined
        assert not any("注：本表为系统依据已核验事实生成" in text for text in texts)


def test_blocks_keep_analysis_only_for_chapters_without_tables(project) -> None:
    """没有系统表格的章节：blocks 只有正文分块。"""
    write_report_section(
        project_id=project["id"], section_key="guarantee", title="担保方案", section_id=f"{project['id']}-guarantee",
        content="### 担保安排\n\n保证方式为保证担保，反担保措施待补材料。",
        content_origin="deterministic", generation_status="deterministic",
    )
    blocks = api_blocks(project, "guarantee")
    assert blocks and not table_blocks(blocks)
    assert any(block["kind"] == "paragraph" and "反担保措施" in (block.get("text") or "") for block in blocks)


def test_missing_table_in_content_still_returns_blocks(project) -> None:
    """历史数据：正文里没有表格块时不报错，表格块标成 present_in_content=false。"""
    write_financial_section(project, with_tables=False)
    try:
        blocks = api_blocks(project)
        produced = table_blocks(blocks)
        assert produced, "表格数据仍应返回，不能因为正文缺表就整章没有表格"
        assert all(block["present_in_content"] is False for block in produced)
        assert all(block["table"]["rows"] for block in produced)
        assert any(block["kind"] == "paragraph" and "负债以短期借款为主" in (block.get("text") or "")
                   for block in blocks)
    finally:
        write_financial_section(project)


def test_exports_keep_table_caption_columns_units_and_notes(project) -> None:
    write_financial_section(project)
    source = chapter_table_blocks(project["id"], "financial", subject=SUBJECT)[0]
    markdown = markdown_report(project["id"], "draft")

    assert source.title in markdown
    assert f"| {source.columns[0].label} |" in markdown and "| --- |" in markdown
    assert "单位：万元" in markdown
    assert source.notes[0] in markdown

    document = Document(io.BytesIO(docx_report(project["id"], "draft")))
    paragraphs = docx_paragraphs(document)
    assert any(source.title in text for text in paragraphs), "Word 必须带表名"
    assert any(source.notes[0] in text for text in paragraphs), "Word 必须带表注"
    assert any("单位：万元" in text for text in paragraphs), "Word 必须带单位与口径"

    labels = [column.label for column in source.columns]
    headers = [[cell.text for cell in table.rows[0].cells] for table in document.tables]
    assert labels in headers, headers
    table = document.tables[headers.index(labels)]
    assert [cell.text for cell in table.rows[1].cells][0] == "资产总计"
    assert [cell.text for cell in table.rows[1].cells][1] == "72,000"

    # 跨页表头重复：首行带 w:tblHeader，长表在第 2 页起仍有列名。
    properties = table.rows[0]._tr.find(qn("w:trPr"))
    assert properties is not None and properties.find(qn("w:tblHeader")) is not None
    # 表名与其表体不分页。
    caption = next(paragraph for paragraph in document.paragraphs if source.title in paragraph.text)
    assert caption.paragraph_format.keep_with_next is True


def _tricky_block():
    """一张含 ``|`` 与换行的系统表格，用于验证转义与原生表格写入。"""
    from backend.app.table_block import annual_series_block, cell_from_text

    return annual_series_block(
        table_id="escape-check", title="转义用例表", subject=SUBJECT, scope="合并口径", unit="万元",
        period_labels=["2025 年末", "2024 年末"],
        rows=[
            ("含|竖线", (cell_from_text("1,000|2,000", unit="万元", source_ids=["s"]),
                        cell_from_text("900", unit="万元", source_ids=["s"]))),
            ("含换行\n科目", (cell_from_text("800", unit="万元", source_ids=["s"]),
                            cell_from_text("未披露"))),
        ],
    )


def test_markdown_export_escapes_cell_pipe_and_newline() -> None:
    """单元格里的竖线与换行必须转义，否则整表会退化成段落或被切成额外列。"""
    from backend.app.report_blocks import escape_markdown_cell, render_table_block_markdown

    block = _tricky_block()
    rendered = render_table_block_markdown(block, index=1, total=1)
    assert "1,000\\|2,000" in rendered, rendered
    assert "含换行<br>科目" in rendered, rendered
    assert escape_markdown_cell("a|b\nc") == "a\\|b<br>c"

    data_line = next(line for line in rendered.splitlines() if "1,000" in line)
    # 反斜杠转义的竖线不再算列分隔符：按「未转义竖线」切分应仍是原列数。
    columns = re.split(r"(?<!\\)\|", data_line)
    assert columns[0].strip() == "" and columns[-1].strip() == ""
    assert len(columns) - 2 == len(block.columns), columns


def test_word_adds_native_table_from_table_block_snapshot() -> None:
    """Word 必须用原生表格 API 从 TableBlock 快照生成，保留全部行且不留竖线文本表格。"""
    from backend.app.exports import _add_table_block, _configure_document
    from backend.app.table_block import table_block_from_dict

    block = table_block_from_dict(_tricky_block().to_dict())
    document = Document()
    _configure_document(document, {"company_name": SUBJECT})
    _add_table_block(document, {"kind": "table", "table": block.to_dict()}, index=1)

    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == [column.label for column in block.columns]
    # 表头 1 行 + 全部数据行，不截断。
    assert len(table.rows) == len(block.rows) + 1
    assert table.rows[1].cells[1].text == "1,000|2,000"
    assert "\n" not in table.rows[2].cells[0].text
    # 跨页重复表头。
    properties = table.rows[0]._tr.find(qn("w:trPr"))
    assert properties is not None and properties.find(qn("w:tblHeader")) is not None
    # 没有任何用带竖线的纯文本冒充表格的段落。
    assert not any(paragraph.text.strip().startswith("|") for paragraph in document.paragraphs)


def test_exports_have_no_forbidden_wording(project) -> None:
    """禁用材料名、内部 UUID、内部字段键都不得进入成品。"""
    write_financial_section(project)
    markdown = markdown_report(project["id"], "draft")
    assert wording_violations(markdown) == [], wording_violations(markdown)
    assert "募集说明书" not in markdown and ORIGINAL_FILENAME not in markdown

    document = Document(io.BytesIO(docx_report(project["id"], "draft")))
    text = "\n".join(docx_paragraphs(document))
    for table in document.tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)
    assert wording_violations(text) == [], wording_violations(text)
    assert ORIGINAL_FILENAME not in text
    # 附录里的来源标签走成品口径（材料名 + 资料编号），不是原件名。
    assert "发行文件〔资料01〕" in text
