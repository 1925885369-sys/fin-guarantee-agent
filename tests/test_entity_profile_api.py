from __future__ import annotations

import json
import io
import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.app.database import connect, rows
from backend.app.main import app
from backend.app.profile_catalog import RECORD_TYPES


def _cleanup(project_id: str = "", profile_ids: list[str] | None = None, import_id: str = "") -> None:
    with connect() as conn:
        if project_id:
            for table in (
                "project_profile_sources", "project_profile_documents", "project_profile_snapshots",
                "project_profile_roles", "project_entities", "projects",
            ):
                conn.execute(f"DELETE FROM {table} WHERE {'id' if table == 'projects' else 'project_id'}=?", (project_id,))
        if import_id:
            conn.execute("DELETE FROM entity_profile_import_files WHERE import_id=?", (import_id,))
            conn.execute("DELETE FROM entity_profile_imports WHERE id=?", (import_id,))
        for profile_id in profile_ids or []:
            for table in (
                "entity_profile_records", "entity_profile_facts", "entity_profile_sources",
                "entity_profile_files", "entity_profile_versions",
            ):
                conn.execute(f"DELETE FROM {table} WHERE profile_id=?", (profile_id,))
            conn.execute("DELETE FROM entity_profiles WHERE id=?", (profile_id,))


def _complete_profile_payload(code: str) -> dict:
    source = {
        "id": "master-source", "filename": "发行人主体档案.pdf", "kind": "主体档案",
        "locator": "第1-20页", "page": 1, "quote": "工商、股权、财务、债务、经营和法律筛查资料汇总。",
    }
    facts = [
        {"fact_key": "legal_representative", "text_value": "张明", "status": "confirmed"},
        {"fact_key": "registered_address", "text_value": "上海市浦东新区金海路1888号", "status": "confirmed"},
        {"fact_key": "registered_capital", "value": 5000, "unit": "万元", "status": "confirmed"},
        {"fact_key": "establishment_date", "text_value": "2012-03-01", "status": "confirmed"},
        {"fact_key": "controller", "text_value": "张明", "status": "confirmed"},
        {"fact_key": "controller_share", "value": 60, "unit": "%", "status": "confirmed"},
        {"fact_key": "interest_bearing_debt", "period": "2025", "value": 1800, "unit": "万元", "scope": "consolidated", "status": "confirmed"},
        {"fact_key": "debt_due_12m", "period": "2025", "value": 600, "unit": "万元", "scope": "consolidated", "status": "confirmed"},
        {"fact_key": "external_guarantee", "period": "2025", "value": 0, "unit": "万元", "scope": "consolidated", "status": "confirmed"},
        {"fact_key": "contingent_liabilities", "period": "2025", "value": 0, "unit": "万元", "scope": "consolidated", "status": "confirmed"},
        {"fact_key": "main_business", "text_value": "精密设备制造", "status": "confirmed"},
        {"fact_key": "top_customer_share", "period": "2025", "value": 18, "unit": "%", "status": "confirmed"},
        {"fact_key": "top_supplier_share", "period": "2025", "value": 12, "unit": "%", "status": "confirmed"},
        {"fact_key": "related_parties", "text_value": "华明科技", "status": "confirmed"},
        {"fact_key": "related_unsupported_amount", "period": "2025", "value": 0, "unit": "万元", "status": "confirmed"},
        {"fact_key": "legal_opinion_status", "text_value": "已完成重大法律事项筛查", "status": "confirmed"},
    ]
    for period in ("2023", "2024", "2025"):
        for fact_key, value in (
            ("total_assets", 8000), ("total_liabilities", 4000), ("revenue", 6000),
            ("equity", 4000), ("net_profit", 500), ("operating_cash_flow", 450),
        ):
            facts.append({
                "fact_key": fact_key, "period": period, "value": value, "unit": "万元",
                "scope": "consolidated", "status": "confirmed",
            })
    for fact in facts:
        fact["source_ids"] = ["master-source"]
    record = {
        "as_of_date": "2026-03-31", "scope": "entity", "status": "confirmed",
        "source_ids": ["master-source"], "data": {"name": "已核验记录"},
    }
    return {
        "name": "完整档案测试有限公司", "unified_code": code,
        "baseline_date": "2025-12-31", "latest_financial_period": "2025",
        "sources": [source], "facts": facts,
        **{record_type: [record] for record_type in RECORD_TYPES},
    }


def test_readiness_and_project_level_transaction_subjects():
    suffix = uuid.uuid4().hex[:10].upper()
    profile_id = project_id = ""
    try:
        with TestClient(app) as client:
            created = client.post("/api/entity-profiles", json=_complete_profile_payload(f"91310000FULL{suffix}"))
            assert created.status_code == 201, created.text
            profile_id = created.json()["id"]
            readiness = client.get(f"/api/entity-profiles/{profile_id}/readiness")
            assert readiness.status_code == 200
            assert readiness.json()["ready_for_agent"] is True
            assert readiness.json()["hard_vetoes"] == []
            assert len(readiness.json()["module_scores"]) == 8

            project = client.post("/api/projects", json={
                "company_name": "将被覆盖", "amount": 20000, "term_months": 36,
                "purpose": "偿还有息债务", "bond_name": "测试债券", "data_consent": True,
                "issuer_profile_id": profile_id, "guarantor_name": "项目保证人有限公司",
                "counter_guarantor_names": ["反担保甲", "反担保乙"],
                "counter_guarantee_status": "exists",
            })
            assert project.status_code == 201, project.text
            project_id = project.json()["id"]
            headers = {"X-Project-Token": project.json()["accessToken"]}
            snapshot = client.get(f"/api/projects/{project_id}/profile-snapshot", headers=headers).json()
            assert [(item["role"], item["profile_id"]) for item in snapshot["roles"]] == [("issuer", profile_id)]
            subjects = rows("SELECT entity_type,name,attributes_json FROM project_entities WHERE project_id=? ORDER BY name", (project_id,))
            assert {(item["entity_type"], item["name"]) for item in subjects} == {
                ("guarantor", "项目保证人有限公司"),
                ("counter_guarantor", "反担保甲"),
                ("counter_guarantor", "反担保乙"),
            }
            evidence = client.get(f"/api/projects/{project_id}/profile-evidence", headers=headers)
            assert evidence.status_code == 200 and evidence.json()["evidence_status"] == "ready"
            assert evidence.json()["documents"] and evidence.json()["sources"]
    finally:
        _cleanup(project_id, [profile_id] if profile_id else [])


def test_profile_rejects_orphan_source_reference():
    payload = {
        "name": "无效来源测试有限公司", "unified_code": f"91310000BAD{uuid.uuid4().hex[:10].upper()}",
        "facts": [{
            "fact_key": "total_assets", "period": "2025", "value": 1000,
            "unit": "万元", "status": "single_source", "source_ids": ["missing-source"],
        }],
    }
    with TestClient(app) as client:
        response = client.post("/api/entity-profiles", json=payload)
        assert response.status_code == 422
        assert "不存在的来源" in response.json()["detail"]


def test_json_import_requires_explicit_selection_and_commits_profile():
    code = f"91310000IMPORT{uuid.uuid4().hex[:8].upper()}"
    payload = {
        "name": "JSON导入测试有限公司", "unified_code": code, "baseline_date": "2025-12-31",
        "sources": [{
            "id": "json-source", "filename": "主体数据.json", "kind": "主体档案",
            "locator": "$.facts[0]", "quote": "2025年资产总额为1000万元。",
        }],
        "facts": [{
            "fact_key": "total_assets", "period": "2025", "value": 1000, "unit": "万元",
            "scope": "consolidated", "status": "confirmed", "source_ids": ["json-source"],
        }],
    }
    profile_id = import_id = ""
    try:
        with TestClient(app) as client:
            created = client.post(
                "/api/entity-profile-imports",
                files={"structured_file": ("profile.json", json.dumps(payload, ensure_ascii=False).encode("utf-8"), "application/json")},
            )
            assert created.status_code == 202, created.text
            import_id = created.json()["id"]
            candidate_ids = [item["id"] for item in created.json()["preview"]["candidates"]]
            empty_commit = client.post(f"/api/entity-profile-imports/{import_id}/commit", json={})
            assert empty_commit.status_code == 422
            committed = client.post(
                f"/api/entity-profile-imports/{import_id}/commit",
                json={"selected_candidate_ids": candidate_ids, "change_note": "测试 JSON 导入"},
            )
            assert committed.status_code == 200, committed.text
            profile_id = committed.json()["profile_id"]
            detail = client.get(f"/api/entity-profiles/{profile_id}")
            assert detail.status_code == 200
            assert detail.json()["facts"][0]["value"] == 1000
            assert detail.json()["change_note"] == "测试 JSON 导入"
    finally:
        _cleanup(profile_ids=[profile_id] if profile_id else [], import_id=import_id)


def test_xlsx_import_builds_candidates_and_commits_profile():
    code = f"91310000XLSX{uuid.uuid4().hex[:8].upper()}"
    workbook = Workbook()
    profile_sheet = workbook.active
    profile_sheet.title = "profile"
    profile_sheet.append(["name", "unified_code", "baseline_date", "latest_financial_period"])
    profile_sheet.append(["Excel导入测试有限公司", code, "2025-12-31", "2025"])
    sources_sheet = workbook.create_sheet("sources")
    sources_sheet.append(["id", "filename", "kind", "locator", "page", "quote"])
    sources_sheet.append(["xlsx-source", "主体档案.xlsx", "主体档案", "facts!A2:H2", 1, "2025年资产总额为1200万元。"])
    facts_sheet = workbook.create_sheet("facts")
    facts_sheet.append(["fact_key", "period", "value", "unit", "scope", "status", "source_ids"])
    facts_sheet.append(["total_assets", "2025", 1200, "万元", "consolidated", "confirmed", "xlsx-source"])
    stream = io.BytesIO()
    workbook.save(stream)
    profile_id = import_id = ""
    try:
        with TestClient(app) as client:
            created = client.post(
                "/api/entity-profile-imports",
                files={
                    "structured_file": (
                        "profile.xlsx", stream.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                },
            )
            assert created.status_code == 202, created.text
            import_id = created.json()["id"]
            candidates = created.json()["preview"]["candidates"]
            assert {item["type"] for item in candidates} == {"profile", "source", "fact"}
            committed = client.post(
                f"/api/entity-profile-imports/{import_id}/commit",
                json={"selected_candidate_ids": [item["id"] for item in candidates]},
            )
            assert committed.status_code == 200, committed.text
            profile_id = committed.json()["profile_id"]
            detail = client.get(f"/api/entity-profiles/{profile_id}").json()
            assert detail["facts"][0]["value"] == 1200
    finally:
        _cleanup(profile_ids=[profile_id] if profile_id else [], import_id=import_id)


def test_frontend_import_contract_targets_profile_and_returns_detail():
    profile_id = import_id = ""
    try:
        with TestClient(app) as client:
            created_profile = client.post("/api/entity-profiles", json={
                "name": "导入契约测试有限公司",
                "unified_code": f"91310000CONTRACT{uuid.uuid4().hex[:6].upper()}",
            })
            assert created_profile.status_code == 201, created_profile.text
            profile_id = created_profile.json()["id"]
            payload = {
                "sources": [{
                    "id": "contract-source", "filename": "财务底稿.json", "kind": "主体档案",
                    "locator": "$.facts[0]", "quote": "2025年资产总额为1000万元。",
                }],
                "facts": [{
                    "fact_key": "total_assets", "period": "2025", "value": 1000, "unit": "万元",
                    "scope": "consolidated", "status": "confirmed", "source_ids": ["contract-source"],
                }],
            }
            started = client.post(
                "/api/entity-profile-imports",
                data={
                    "target_profile_id": profile_id, "package_kind": "financial",
                    "baseline_date": "2025-12-31", "period": "2025", "scope": "consolidated",
                },
                files={"structured_file": ("profile.json", json.dumps(payload).encode(), "application/json")},
            )
            assert started.status_code == 202, started.text
            job = started.json()
            import_id = job["id"]
            assert job["target_profile_id"] == profile_id
            assert job["status"] == "ready_to_commit"
            assert job["package_kind"] == "financial"
            assert job["files"][0]["status"] == "parsed"
            assert len(job["candidates"]) == 1
            candidate = job["candidates"][0]
            assert candidate["fact_key"] == "total_assets"
            assert candidate["source_ids"] == ["contract-source"]

            committed = client.post(f"/api/entity-profile-imports/{import_id}/commit", json={
                "selected_candidate_ids": [candidate["id"]],
                "corrections": [{"candidate_id": candidate["id"], "value": 1200}],
                "change_summary": "前端字段级修正",
            })
            assert committed.status_code == 200, committed.text
            detail = committed.json()
            assert detail["id"] == profile_id
            assert detail["profile_id"] == profile_id
            assert detail["version"] == 2
            assert detail["facts"][0]["value"] == 1200
            assert detail["change_note"] == "前端字段级修正"
    finally:
        _cleanup(profile_ids=[profile_id] if profile_id else [], import_id=import_id)


def test_frozen_incomplete_snapshot_blocks_run_and_report_generation():
    profile_id = project_id = ""
    try:
        with TestClient(app) as client:
            profile = client.post("/api/entity-profiles", json={
                "name": "门禁测试发行人有限公司",
                "unified_code": f"91310000GATE{uuid.uuid4().hex[:8].upper()}",
            })
            assert profile.status_code == 201, profile.text
            profile_id = profile.json()["id"]
            readiness = client.get(f"/api/entity-profiles/{profile_id}/readiness").json()
            assert readiness["ready_for_agent"] is False
            assert all({"key", "label", "reason"} <= set(item) for item in readiness["hard_vetoes"])
            assert all("status" in item for item in readiness["module_scores"])
            assert all("missing_fact_keys" in item for item in readiness["requirements"])

            project = client.post("/api/projects", json={
                "company_name": "将被覆盖", "amount": 1000, "term_months": 36,
                "purpose": "债券发行", "bond_name": "门禁测试债券", "data_consent": True,
                "issuer_profile_id": profile_id,
            })
            assert project.status_code == 201, project.text
            project_id = project.json()["id"]
            headers = {"X-Project-Token": project.json()["accessToken"]}
            for path in (f"/api/projects/{project_id}/runs", f"/api/projects/{project_id}/report/generate"):
                blocked = client.post(path, json={"focus": "测试门禁"}, headers=headers)
                assert blocked.status_code == 409
                detail = blocked.json()["detail"]
                assert detail["code"] == "profile_snapshot_not_ready"
                assert detail["hard_vetoes"]
                assert detail["snapshot_id"]
    finally:
        _cleanup(project_id, [profile_id] if profile_id else [])
