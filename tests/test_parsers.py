from pathlib import Path

import fitz
import openpyxl
import pytest
from docx import Document as WordDocument

from backend.app.facts import deterministic_extract, normalize_numeric
from backend.app.parsers import parse_document


def test_supported_document_formats_are_parseable(tmp_path: Path):
    """三种受支持格式都要能解析。

    产品已收敛为仅债券担保：债券案例材料全部是 PDF，docx/xlsx 只出现在用户上传侧，
    因此这里直接构造三种格式的样本，不再依赖已删除的贷款演示材料生成器。
    """
    docx_path = tmp_path / "sample.docx"
    document = WordDocument()
    document.add_heading("发行人基本情况", level=1)
    for index, line in enumerate([
        "发行人主营业务为精密零部件制造，2025 年营业收入 48000 万元，净利润 2800 万元。",
        "报告期内营业收入分别为 2023 年 39360 万元、2024 年 43680 万元、2025 年 48000 万元。",
        "经营活动现金流净额分别为 4680 万元、5590 万元、6500 万元。",
        "应收账款账面价值 9840 万元，已计提坏账准备 320 万元。",
        "有息负债合计 25420 万元，其中短期借款 9151 万元、长期借款 11693 万元。",
        "本期债券拟发行规模 20000 万元，期限 3 年，用于补充流动资金。",
        "募集资金专户由受托管理人监管，用途变更需履行约定程序。",
        "评级机构认为主体信用等级 AA-，评级展望稳定。",
    ], start=1):
        document.add_paragraph(f"{index}. {line}")
    document.save(docx_path)

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["科目", "2023", "2024", "2025"])
    for row in (
        ["营业收入", 39360, 43680, 48000],
        ["营业成本", 29120, 32320, 35520],
        ["净利润", 2296, 2548, 2800],
        ["经营活动现金流净额", 4680, 5590, 6500],
        ["应收账款", 8520, 9180, 9840],
        ["存货", 4350, 4680, 5020],
        ["资产总额", 62400, 66800, 71200],
        ["负债总额", 43680, 46760, 49840],
        ["有息负债", 22400, 23900, 25420],
    ):
        sheet.append(row)
    workbook.save(xlsx_path)

    pdf_path = tmp_path / "sample.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        lines = [
            "Bond guarantee due diligence sample",
            "Issuer main business: precision components manufacturing.",
            "Revenue 2023/2024/2025: 39360 / 43680 / 48000 wan yuan.",
            "Net profit 2023/2024/2025: 2296 / 2548 / 2800 wan yuan.",
            "Operating cash flow: 4680 / 5590 / 6500 wan yuan.",
            "Interest bearing debt 2025: 25420 wan yuan.",
            "Proposed bond issue amount: 20000 wan yuan, term 3 years.",
            "Use of proceeds: working capital replenishment.",
            "Rating agency opinion: AA- with stable outlook.",
        ]
        for index, line in enumerate(lines):
            page.insert_text((72, 90 + index * 18), line)
        pdf.save(pdf_path)

    for path in (docx_path, xlsx_path, pdf_path):
        parsed = parse_document(path)
        assert parsed.fragments, path.name
        assert parsed.quality > 0.2, path.name


def test_scanned_pdf_is_flagged(tmp_path: Path):
    path = tmp_path / "scan.pdf"
    document = fitz.open(); document.new_page(); document.save(path)
    parsed = parse_document(path)
    assert parsed.quality <= 0.2
    assert "扫描型" in parsed.warning


def test_bond_prospectus_over_150_pages_is_parseable(tmp_path: Path):
    path = tmp_path / "prospectus.pdf"
    document = fitz.open()
    for page_number in range(151):
        page = document.new_page()
        page.insert_text(
            (50, 50),
            f"prospectus page {page_number + 1} issuance terms, use of proceeds, repayment arrangements, "
            "investor protection, guarantee scope and material risk disclosures",
        )
    document.save(path)

    parsed = parse_document(path)

    assert parsed.page_count == 151
    assert len(parsed.fragments) == 151
    assert parsed.quality > 0.2


def test_encrypted_and_corrupt_files_fail_cleanly(tmp_path: Path):
    encrypted = tmp_path / "encrypted.pdf"
    document = fitz.open(); page = document.new_page(); page.insert_text((50, 50), "secret")
    document.save(encrypted, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    with pytest.raises(ValueError, match="加密"):
        parse_document(encrypted)
    corrupt = tmp_path / "broken.xlsx"; corrupt.write_bytes(b"PK broken")
    with pytest.raises(ValueError, match="无法打开"):
        parse_document(corrupt)


def test_unit_normalization_and_labelled_excel_extraction():
    assert normalize_numeric(1.2, "亿元", "amount") == (12000.0, "万元")
    assert normalize_numeric(50000, "元", "amount") == (5.0, "万元")
    sources = [
        {"id": "h", "document_id": "d", "locator_type": "xlsx_range", "locator": "财务!A1:C1", "text": "指标（万元） | 2024 | 2025"},
        {"id": "r", "document_id": "d", "locator_type": "xlsx_range", "locator": "财务!A2:C2", "text": "营业收入 | 1000 | 1200"},
    ]
    facts = deterministic_extract(sources)
    assert [(item["period"], item["value"]) for item in facts if item["fact_key"] == "revenue"] == [("2024", 1000.0), ("2025", 1200.0)]


def test_excel_amount_that_looks_like_a_year_uses_header_periods():
    sources = [
        {"id": "h", "document_id": "d", "locator_type": "xlsx_range", "locator": "财务!A1:C1", "text": "指标（万元） | 2024 | 2025"},
        {"id": "r", "document_id": "d", "locator_type": "xlsx_range", "locator": "财务!A2:C2", "text": "流动资产 | 1700 | 2000"},
    ]
    facts = deterministic_extract(sources)

    assert [(item["period"], item["value"]) for item in facts if item["fact_key"] == "current_assets"] == [
        ("2024", 1700.0),
        ("2025", 2000.0),
    ]


def test_compact_paired_year_prose_extracts_each_period_and_short_labels():
    sources = [{
        "id": "compact", "document_id": "d", "locator_type": "xlsx_range", "locator": "财务!A2:C10",
        "text": "2024/2025营业收入2600/3300；净利润320/400；经营现金流210/300；应收700/850；资产3000；负债1200。",
    }]

    facts = deterministic_extract(sources)

    def values(key: str) -> list[tuple[str, float]]:
        return [(item["period"], item["value"]) for item in facts if item["fact_key"] == key]

    assert values("revenue") == [("2024", 2600), ("2025", 3300)]
    assert values("net_profit") == [("2024", 320), ("2025", 400)]
    assert values("accounts_receivable") == [("2024", 700), ("2025", 850)]
    assert values("total_assets") == [("2025", 3000)]
    assert values("total_liabilities") == [("2025", 1200)]


def test_short_aliases_do_not_map_customer_share_or_specific_liabilities_to_generic_fields():
    sources = [{
        "id": "mixed", "document_id": "d", "locator_type": "docx_paragraph", "locator": "第1段",
        "text": "2025年第一大客户收入占比48%，有息负债700万元，流动负债800万元，净资产1800万元。",
    }]

    facts = deterministic_extract(sources)

    assert not [item for item in facts if item["fact_key"] == "revenue"]
    assert not [item for item in facts if item["fact_key"] == "total_liabilities"]
    assert not [item for item in facts if item["fact_key"] == "total_assets"]
    assert next(item for item in facts if item["fact_key"] == "top_customer_share")["value"] == 48
    assert next(item for item in facts if item["fact_key"] == "interest_bearing_debt")["value"] == 700


def test_multi_field_labelled_text_uses_the_number_next_to_each_label():
    sources = [{
        "id": "financials", "document_id": "d", "locator_type": "xlsx_range", "locator": "年度!A2:J2",
        "text": "年度2025 | 营业收入9500万元 | 净利润280万元 | 经营活动现金流净额-100万元 | 应收账款1900万元 | 资产总额6000万元 | 负债总额4500万元 | 流动资产4000万元 | 流动负债3500万元 | 有息负债2200万元",
    }]

    facts = {item["fact_key"]: item for item in deterministic_extract(sources)}

    assert facts["revenue"]["value"] == 9500
    assert facts["net_profit"]["value"] == 280
    assert facts["operating_cash_flow"]["value"] == -100
    assert facts["accounts_receivable"]["value"] == 1900
    assert facts["equity"]["value"] == 1500


def test_negative_numeric_statement_extracts_zero_without_borrowing_another_field_value():
    sources = [{
        "id": "risk", "document_id": "d", "locator_type": "docx_paragraph", "locator": "第1段",
        "text": "2025年利息支出120万元；无对外担保及其他或有负债。",
    }]

    facts = {item["fact_key"]: item for item in deterministic_extract(sources)}

    assert facts["interest_expense"]["value"] == 120
    assert facts["external_guarantee"]["value"] == 0
    assert facts["contingent_liabilities"]["value"] == 0


def test_unrelated_numbers_are_not_used_for_a_label_without_a_local_value():
    sources = [{
        "id": "governance", "document_id": "d", "locator_type": "docx_paragraph", "locator": "第1段",
        "text": "实际控制人持股65%；第一还款来源为销售回款。",
    }]

    facts = deterministic_extract(sources)

    assert not [item for item in facts if item["fact_key"] == "operating_inflow"]


def test_common_prose_financial_labels_are_extracted_without_false_asset_match():
    sources = [{
        "id": "prose", "document_id": "d", "locator_type": "docx_paragraph", "locator": "第4段",
        "text": "2025年资产5200万元、负债3500万元、净资产1700万元、到期债务900万元。纳税收入2750万元，银行经营流入2850万元；用途合同450万元；估值距今18个月。",
    }]

    facts = {item["fact_key"]: item for item in deterministic_extract(sources)}

    assert facts["total_assets"]["value"] == 5200
    assert facts["total_liabilities"]["value"] == 3500
    assert facts["equity"]["value"] == 1700
    assert facts["debt_due_12m"]["value"] == 900
    assert facts["tax_revenue"]["value"] == 2750
    assert facts["operating_inflow"]["value"] == 2850
    assert facts["supporting_contract_amount"]["value"] == 450
    assert facts["collateral_valuation_age_months"]["value"] == 18


def test_overdue_receivables_are_derived_from_aging_buckets():
    sources = [
        {"id": "one", "document_id": "d", "locator_type": "xlsx_range", "locator": "账龄!A2:C2", "text": "1年以内 | 1100 | 66.7%"},
        {"id": "two", "document_id": "d", "locator_type": "xlsx_range", "locator": "账龄!A3:C3", "text": "1-2年 | 350 | 21.2%"},
        {"id": "three", "document_id": "d", "locator_type": "xlsx_range", "locator": "账龄!A4:C4", "text": "2年以上 | 200 | 12.1%"},
    ]

    facts = deterministic_extract(sources)
    overdue = next(item for item in facts if item["fact_key"] == "overdue_receivables")

    assert overdue["value"] == 550
    assert overdue["period"] == ""
    assert overdue["source_ids"] == ["two", "three"]
